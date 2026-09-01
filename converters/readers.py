"""Readers: platform export -> :class:`Product`.

One function per source format, all returning the same model. Nothing here
knows about any output format -- see :mod:`converters.product` for why that
separation is enforced.

Format detection is by header content rather than filename, because these files
arrive named ``export (3).csv`` and the extension says nothing about which
platform produced them.
"""

from __future__ import annotations

import csv
import os

from .product import Product, Variant

#: Any spelling of a photo/image column, mapped to its slot number.
_PHOTO_PREFIXES = ("photo", "image")


def _photo_slot(name: str) -> int | None:
    cleaned = (name or "").strip().lower().replace("_", " ")
    for prefix in _PHOTO_PREFIXES:
        if cleaned.startswith(prefix):
            digits = cleaned[len(prefix):].strip()
            if digits.isdigit():
                return int(digits)
    return None


def photo_fields(header: list[str]) -> list[str]:
    """Photo columns in slot order, whatever they are spelled."""
    numbered = [(slot, name) for name in header
                if (slot := _photo_slot(name)) is not None]
    return [name for _, name in sorted(numbered)]


def _get(row: dict, *names: str) -> str:
    for name in names:
        value = (row.get(name) or "").strip()
        if value:
            return value
    return ""


def detect_format(header: list[str]) -> str:
    """``"altera"``, ``"vela"`` or ``"unknown"``, from the header alone."""
    lowered = {h.strip().lower() for h in header}
    if "handle" in lowered and ("image src" in lowered or "variant sku" in lowered):
        return "altera"
    if "title" in lowered and (photo_fields(header) or "sku" in lowered):
        return "vela"
    return "unknown"


def read_vela_rows(rows, header: list[str]) -> list[Product]:
    """Vela/Etsy CSV: parent rows with variation rows beneath them."""
    columns = photo_fields(header)
    products: list[Product] = []
    current: Product | None = None

    for row in rows:
        title = _get(row, "TITLE", "Title", "title")
        if title:
            current = Product(
                title=title,
                description=_get(row, "DESCRIPTION", "Description"),
                sku=_get(row, "SKU", "Sku", "sku"),
                price=_get(row, "PRICE", "Price"),
                currency=_get(row, "CURRENCY_CODE", "Currency"),
                quantity=_get(row, "QUANTITY", "Quantity"),
                tags=[t.strip() for t in _get(row, "TAGS", "Tags").split(",") if t.strip()],
            )
            products.append(current)
        elif current is None:
            continue

        for column in columns:
            url = (row.get(column) or "").strip()
            if url and url not in current.images:
                current.images.append(url)

        value = _get(row, "VARIATION 1 VALUES", "Var Value 1")
        if value:
            current.variants.append(Variant(
                option_name=_get(row, "VARIATION 1 NAME"),
                option_value=value,
                price=_get(row, "Var Price"),
                sku=_get(row, "Var SKU"),
            ))
    return products


def read_altera_rows(rows) -> list[Product]:
    """Altera/Shopify CSV: handle-grouped blocks."""
    products: dict[str, Product] = {}
    order: list[str] = []
    current: str | None = None

    for row in rows:
        handle = _get(row, "Handle", "handle")
        if handle:
            current = handle
            if handle not in products:
                products[handle] = Product(title=_get(row, "Title"))
                order.append(handle)
            product = products[handle]
            if not product.title:
                product.title = _get(row, "Title")
            product.description = product.description or _get(row, "Body HTML")
            product.vendor = product.vendor or _get(row, "Vendor")
            if not product.tags:
                tags = _get(row, "Tags")
                product.tags = [t.strip() for t in tags.split(",") if t.strip()]
        elif current is None:
            continue

        product = products[current]
        image = _get(row, "Image Src")
        if image and image not in product.images:
            product.images.append(image)

        sku = _get(row, "Variant SKU")
        if sku:
            if not product.sku:
                product.sku = sku
            product.variants.append(Variant(
                option_name=_get(row, "Option1 Name"),
                option_value=_get(row, "Option1 Value"),
                price=_get(row, "Variant Price"),
                sku=sku,
            ))
    return [products[h] for h in order]


def read_catalogue(path: str) -> list[Product]:
    """Read any supported export into Products, sniffing the format."""
    extension = os.path.splitext(path)[1].lower()
    if extension in (".xlsx", ".xlsm"):
        return _read_xlsx(path)

    # utf-8-sig: both platforms hand out BOM-prefixed CSVs, and a BOM welded to
    # the first header name silently breaks TITLE/Handle lookup.
    with open(path, newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        header = list(reader.fieldnames or [])
        rows = list(reader)

    kind = detect_format(header)
    if kind == "altera":
        return read_altera_rows(rows)
    if kind == "vela":
        return read_vela_rows(rows, header)
    raise ValueError(
        "Could not identify %s as a Vela/Etsy or Altera/Shopify export. "
        "Header was: %s" % (path, ", ".join(header[:8]))
    )


def _read_xlsx(path: str) -> list[Product]:
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    try:
        name = next((n for n in book.sheetnames
                     if n.strip().lower() == "products"), book.sheetnames[0])
        rows = book[name].iter_rows(values_only=True)
        try:
            header = [str(c).strip() if c is not None else "" for c in next(rows)]
        except StopIteration:
            return []
        dicts = [{h: ("" if v is None else str(v)) for h, v in zip(header, row)}
                 for row in rows]
    finally:
        book.close()

    kind = detect_format(header)
    if kind == "altera":
        return read_altera_rows(dicts)
    if kind == "vela":
        return read_vela_rows(dicts, header)
    raise ValueError("Could not identify the sheet in %s" % path)
